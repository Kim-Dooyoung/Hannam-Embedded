from openpyxl import Workbook, load_workbook
import os

file_name = "임베디드_음식.xlsx"

if os.path.exists(file_name):
    # 파일이 존재하면 불러오기
    wb = load_workbook(file_name)
    ws = wb.active
else:
    # 파일이 없으면 새로 생성
    wb = Workbook()
    ws = wb.active
    ws.title = "음식"

    column = ['번호', '음식']
    ws.append(column)

# 재료 입력 (공백으로 구분된 여러 음식 입력 가능)
Food = input("궁금한 음식을 입력하세요(입력 후 엔터 클릭): ")

# 입력된 음식을 공백을 기준으로 나누어 리스트로 만듭니다.
food_list = Food.split()

existing_rows = ws.max_row

for idx, food in enumerate(food_list, start=existing_rows):
    ws.append([idx, food])

# 엑셀 파일 저장
wb.save(file_name)
